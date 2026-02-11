"""文件解析工具 - 支持多种文档格式"""
import os
from pathlib import Path

import chardet


def detect_encoding(file_path: str) -> str:
    with open(file_path, "rb") as f:
        raw = f.read(10000)
    result = chardet.detect(raw)
    return result.get("encoding") or "utf-8"


def parse_txt(file_path: str) -> str:
    encoding = detect_encoding(file_path)
    with open(file_path, "r", encoding=encoding, errors="replace") as f:
        return f.read()


def parse_md(file_path: str) -> str:
    return parse_txt(file_path)


def parse_pdf(file_path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n\n".join(texts)


def parse_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    texts = []
    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                texts.append(row_text)
    return "\n\n".join(texts)


def parse_csv(file_path: str) -> str:
    import csv
    encoding = detect_encoding(file_path)
    texts = []
    with open(file_path, "r", encoding=encoding, errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            texts.append(" | ".join(row))
    return "\n".join(texts)


def parse_html(file_path: str) -> str:
    from bs4 import BeautifulSoup
    encoding = detect_encoding(file_path)
    with open(file_path, "r", encoding=encoding, errors="replace") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


PARSERS = {
    ".txt": parse_txt,
    ".md": parse_md,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".doc": parse_docx,
    ".csv": parse_csv,
    ".html": parse_html,
    ".htm": parse_html,
}


def parse_file(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"不支持的文件格式: {ext}")
    return parser(file_path)


def get_file_size(file_path: str) -> int:
    return os.path.getsize(file_path)
